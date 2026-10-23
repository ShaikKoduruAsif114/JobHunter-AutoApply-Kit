#!/usr/bin/env python3
"""
Submission Tracker & Excel Synchronizer
Detects confirmed job application submissions directly from browser history,
updates pipeline state, and logs entries into an Excel workbook.
"""

import os
import re
import json
import shutil
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Any, List

def sync_submissions(chrome_user_data_dir: Path, state_file: Path, excel_path: Path) -> int:
    """Reads browser history, matches submission confirmation URLs, and updates state + Excel."""
    history_db = chrome_user_data_dir / "Default" / "History"
    if not history_db.exists():
        return 0

    tmp_db = Path("cache_history_tmp.db")
    try:
        shutil.copy2(history_db, tmp_db)
        conn = sqlite3.connect(tmp_db)
        cursor = conn.cursor()
        cursor.execute('''
            SELECT url, title, last_visit_time 
            FROM urls 
            WHERE url LIKE '%confirmation%' 
               OR url LIKE '%thankyou%'
               OR url LIKE '%success%'
               OR url LIKE '%/thanks%'
               OR title LIKE '%thank you%' 
               OR title LIKE '%received%' 
            ORDER BY last_visit_time DESC
        ''')
        rows = cursor.fetchall()
        conn.close()
    except Exception:
        rows = []
    finally:
        if tmp_db.exists():
            try:
                os.remove(tmp_db)
            except Exception:
                pass

    if not rows:
        return 0

    # Load State
    state = {"jobs": {}}
    if state_file.exists():
        try:
            with open(state_file, "r", encoding="utf-8") as f:
                state = json.load(f)
        except Exception:
            state = {"jobs": {}}

    jobs = state.get("jobs", {})
    epoch_start = datetime(1601, 1, 1, tzinfo=timezone.utc)
    new_submissions = []

    for u, t, lvt in rows:
        ts = epoch_start.timestamp() + (lvt / 1000000.0)
        sub_dt = datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S")

        matched = None
        for jid, j in jobs.items():
            apply_url = j.get("apply_link", "")
            base_clean = apply_url.split("?")[0].rstrip("/")
            if base_clean and base_clean in u:
                matched = j
                break
            # Match greenhouse ID
            m = re.search(r'/jobs/(\d+)', apply_url)
            if m and m.group(1) in u:
                matched = j
                break
            # Match lever token
            m_lev = re.search(r'([a-f0-9]{8}-[a-f0-9]{4}-[a-f0-9]{4}-[a-f0-9]{4}-[a-f0-9]{12})', apply_url)
            if m_lev and m_lev.group(1) in u:
                matched = j
                break

        if matched and matched.get("status") != "submitted":
            matched["status"] = "submitted"
            matched["submitted_at"] = sub_dt
            new_submissions.append(matched)

    if new_submissions:
        with open(state_file, "w", encoding="utf-8") as f:
            json.dump(state, f, indent=2)

        # Sync to Excel
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment

            if excel_path.exists():
                wb = openpyxl.load_workbook(excel_path)
                sheet = wb.active
            else:
                wb = openpyxl.Workbook()
                sheet = wb.active
                sheet.title = "Submitted Applications"
                headers = ["Company", "Role Title", "Location", "Track", "ATS Platform", "Apply Link", "Status", "Submitted At", "Notes"]
                sheet.append(headers)
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center")

            existing_links = set(sheet.cell(row=r, column=6).value.strip().lower() for r in range(2, sheet.max_row + 1) if sheet.cell(row=r, column=6).value)

            for sub in new_submissions:
                clean_link = (sub.get("apply_link") or "").strip().lower()
                if clean_link not in existing_links:
                    sheet.append([
                        sub.get("company", ""),
                        sub.get("title", ""),
                        sub.get("location", ""),
                        sub.get("track", "SDE"),
                        sub.get("ats", "Direct"),
                        sub.get("apply_link", ""),
                        "Submitted",
                        sub.get("submitted_at", ""),
                        "Confirmed in browser"
                    ])
                    existing_links.add(clean_link)

            wb.save(excel_path)
        except Exception as e:
            print(f"⚠️ Excel update notice: {e}")

    return len(new_submissions)
